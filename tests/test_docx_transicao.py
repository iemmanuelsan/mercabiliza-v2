"""Testes do formulário DOCX de transição contábil.

O teste que mais importa aqui é o de **cobertura da planilha**: o documento
substitui um Excel que a equipe usa em produção, e perder um campo na troca
significaria descobrir a falta no meio de um onboarding, com o cliente na
frente. A lista de rótulos abaixo é transcrita da planilha anexada pelo Iago.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from src.core.models import Endereco
from src.core.pessoas import ContratantePF, ContratantePJ, RepresentanteLegal
from src.exporters.docx_transicao import (
    DOCUMENTOS_A_SOLICITAR,
    CampoDesconhecidoError,
    MARCADOR_PENDENTE,
    dados_de_contratante,
    gerar_formulario_transicao,
)

END = Endereco(logradouro="RUA DAS FLORES", numero="88", bairro="Centro",
               municipio="Jundiaí", uf="SP", cep="13201000")

PJ = ContratantePJ(
    razao_social="SUZANA DAS DORES MARCILIA COMERCIO DE ALIMENTOS LTDA",
    cnpj="32348868000173",
    nome_fantasia="MiniMercado Autônomo",
    regime="Simples Nacional",
    inscricao_estadual="159384180119",
    cnae_principal="4712100 - Minimercados",
    endereco=END,
    telefone="11 95076-0336",
    email="suzanamarcilia@yahoo.com.br",
    representante=RepresentanteLegal(
        nome="Suzana das Dores Marcilia", cpf="52998224725"),
)

PF = ContratantePF(nome="Vinicius Almeida", cpf="11144477735", endereco=END,
                   telefone="(22) 98888-7777", email="v@exemplo.com.br")


def _texto(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                partes.append(celula.text)
    return "\n".join(partes)


def _runs(docx_bytes: bytes):
    """Todos os runs do documento, para inspecionar negrito e cor."""
    doc = Document(io.BytesIO(docx_bytes))
    saida = []
    for p in doc.paragraphs:
        saida.extend(p.runs)
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for p in celula.paragraphs:
                    saida.extend(p.runs)
    return saida


# --------------------------------------------------------------------------- #
# Cobertura da planilha que este documento substitui                           #
# --------------------------------------------------------------------------- #
# Transcritos da planilha "FORMULÁRIO DE ENTRADA DE NOVOS CLIENTES". Onde o
# rótulo foi reescrito para caber ou para ficar mais claro, comparo por um
# trecho estável em vez do texto integral.
CAMPOS_DA_PLANILHA = [
    # Informações iniciais
    "Modelo de entrada",
    "Competência de entrada",
    "Razão social",
    "CNPJ",
    "Tem filiais?",
    "Segmento da empresa",
    "Faturamento mensal médio",
    "CNPJ das filiais",
    "Quantos sócios?",
    "Responsável legal",
    "CPF do responsável",
    "E-mail 1",
    "E-mail 2",
    "Telefone 1",
    "Telefone 2",
    # Departamento pessoal
    "Tem funcionários?",
    "Se sim, quantos?",
    "Terá pró-labore?",
    "Terá adiantamento salarial?",
    # Departamento fiscal
    "Possui certificado digital válido?",
    "Sistema de notas utilizado",
    "Tipo de empresa",
    "Regime tributário",
    # Sucessão contábil
    "Nome do escritório / contador",
    "E-mail",
    "Telefone",
]

# ------------------------------------------------------------------------- #
# O que foi DELIBERADAMENTE removido (revisão de 09/09/2026)                 #
# ------------------------------------------------------------------------- #
# O Iago cortou estes blocos e campos depois de ver o template. Guardar a
# lista aqui, como asserção de AUSÊNCIA, tem duas funções: registra que a
# falta é decisão e não esquecimento, e impede que voltem por engano numa
# edição futura — o teste de cobertura acima sozinho não pegaria isso, porque
# ele só verifica presença.
#
# Motivo de cada corte:
#   Onboarding e comunicação → definido internamente, não pelo cliente
#   Departamento contábil    → idem
#   Fator R e monofásicos    → apuração nossa, não pergunta ao cliente
#   Departamento financeiro  → honorários e 13º ficam no contrato
#   Observações finais       → conversa, não formulário
#   CRC / último mês / débito em aberto → tratativa entre contadores
#   Assinatura, local e data → o documento deixou de ser termo assinado
#   Contatos da Mercabiliza  → os valores em config.py são de outra empresa
REMOVIDOS_DE_PROPOSITO = [
    "Onboarding",
    "Será criado grupo de WhatsApp?",
    "Tipo de onboarding",
    "Departamento contábil",
    "Tipo de balanço",
    "Aplica-se Fator R?",
    "monofásicos",
    "Departamento financeiro",
    "Valor de honorários",
    "Terá 13º",
    "Contrata todos os departamentos?",
    "Observações finais",
    "Qual o motivo da troca de contabilidade?",
    "Observações adicionais",
    "CRC do contador anterior",
    "Último mês escriturado",
    "honorários em aberto",
    "Assinatura",
    "Local e data",
]


@pytest.mark.parametrize("rotulo", CAMPOS_DA_PLANILHA)
def test_formulario_cobre_todos_os_campos_da_planilha(rotulo: str) -> None:
    """Nenhum campo mantido da planilha pode sumir na migração para DOCX."""
    assert rotulo in _texto(gerar_formulario_transicao())


@pytest.mark.parametrize("trecho", REMOVIDOS_DE_PROPOSITO)
def test_o_que_foi_removido_nao_volta(trecho: str) -> None:
    assert trecho not in _texto(gerar_formulario_transicao())


# --------------------------------------------------------------------------- #
# Comportamento em branco                                                      #
# --------------------------------------------------------------------------- #
def test_sem_dados_gera_formulario_em_branco_utilizavel() -> None:
    """Sem argumento nenhum o documento tem de sair — é o equivalente ao
    impresso que a equipe usa hoje."""
    texto = _texto(gerar_formulario_transicao())
    assert "FORMULÁRIO DE ENTRADA DE NOVOS CLIENTES" in texto
    assert MARCADOR_PENDENTE in texto


def test_modelo_de_entrada_ja_vem_preenchido() -> None:
    # É a única resposta que não varia: este formulário é, por definição, de
    # transição contábil. Deixá-lo em branco convidaria a errar o próprio
    # propósito do documento.
    assert "Transição contábil" in _texto(gerar_formulario_transicao())


# --------------------------------------------------------------------------- #
# A convenção visual — que é o motivo de o documento existir                   #
# --------------------------------------------------------------------------- #
def test_dado_do_sistema_sai_em_negrito_e_sem_marcador() -> None:
    docx = gerar_formulario_transicao(
        dados_iniciais={"razao_social": "MERCADO DO ZE LTDA"})

    negritos = [r.text for r in _runs(docx) if r.bold]
    assert "MERCADO DO ZE LTDA" in negritos

    # O campo preenchido não pode continuar pedindo preenchimento.
    doc = Document(io.BytesIO(docx))
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                if "Razão social" in celula.text:
                    assert MARCADOR_PENDENTE not in celula.text


def test_campo_ausente_continua_pedindo_preenchimento() -> None:
    docx = gerar_formulario_transicao(
        dados_iniciais={"razao_social": "MERCADO DO ZE LTDA"})
    doc = Document(io.BytesIO(docx))
    encontrou = False
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                if "Competência de entrada" in celula.text:
                    assert MARCADOR_PENDENTE in celula.text
                    encontrou = True
    assert encontrou, "a seção de informações iniciais não foi renderizada"


# --------------------------------------------------------------------------- #
# Sucessão contábil                                                            #
# --------------------------------------------------------------------------- #
def test_lista_do_contador_anterior_esta_no_documento() -> None:
    texto = _texto(gerar_formulario_transicao())
    for grupo, itens in DOCUMENTOS_A_SOLICITAR.items():
        assert grupo in texto
        for item in itens:
            assert item in texto, f"faltou '{item}' no grupo {grupo}"


def test_lista_e_informativa_e_diz_que_o_cliente_nao_precisa_agir() -> None:
    """O Iago pediu explicitamente que o cliente saiba que QUEM SOLICITA é a
    Mercabiliza. Sem essa frase o cliente lê a lista como tarefa dele e a
    transição trava esperando um movimento que nunca vem."""
    texto = _texto(gerar_formulario_transicao())
    assert "você não precisa providenciar nada disto" in texto
    assert "nós mesmos entramos em contato" in texto


def test_declaracao_informa_sem_prometer_assinatura() -> None:
    """A declaração ficou, a assinatura saiu — e as duas coisas precisam ser
    coerentes.

    A versão anterior dizia "DECLARO... e AUTORIZO a Mercabiliza a
    solicitar...". Num documento sem linha de assinatura não há quem declare
    nem quem autorize, e afirmar que a autorização foi concedida é pior do que
    não ter o texto: o contador anterior pode recusar a entrega, e a transição
    trava exatamente onde este formulário deveria destravar.

    A autorização real está no contrato de prestação de serviços, que é
    assinado. Aqui o texto apenas INFORMA.
    """
    texto = _texto(gerar_formulario_transicao())
    assert "LGPD" in texto
    assert "contrato de prestação de serviços" in texto
    # Nada pode afirmar que este documento concede a autorização.
    assert "autorizo" not in texto.lower()
    assert "declaro" not in texto.lower()


def test_documento_nao_traz_contato_da_mercabiliza() -> None:
    """O telefone e o e-mail em ``config.py`` são de outra contabilidade.
    Imprimir contato errado é pior do que não imprimir contato nenhum."""
    from src.config import CONTRATADA_EMAIL, CONTRATADA_TELEFONE

    texto = _texto(gerar_formulario_transicao())
    assert CONTRATADA_EMAIL not in texto
    assert CONTRATADA_TELEFONE not in texto


def test_contato_da_contabilidade_anterior_e_campo_de_formulario() -> None:
    docx = gerar_formulario_transicao(
        dados_sucessao={"email_anterior": "contato@agilize.com.br"})
    negritos = [r.text for r in _runs(docx) if r.bold]
    assert "contato@agilize.com.br" in negritos


# --------------------------------------------------------------------------- #
# Prefill a partir do contratante consultado                                   #
# --------------------------------------------------------------------------- #
def test_prefill_de_pj_traz_cadastro_formatado() -> None:
    dados = dados_de_contratante(PJ)
    assert dados["razao_social"] == PJ.razao_social
    # Formatado, não cru: o cliente confere contra o cartão CNPJ.
    assert dados["cnpj"] == "32.348.868/0001-73"
    assert dados["cpf_responsavel"] == "529.982.247-25"
    assert dados["inscricao_estadual"] == "159384180119"
    assert dados["regime"] == "Simples Nacional"


def test_prefill_nao_inventa_segmento_a_partir_do_cnae() -> None:
    """O CNAE ("comércio varejista de mercadorias em geral") não é o segmento
    ("minimercado autônomo"). Preencher um com o outro pareceria útil e
    colocaria no documento uma resposta que ninguém conferiu."""
    dados = dados_de_contratante(PJ)
    assert "segmento" not in dados

    docx = gerar_formulario_transicao(dados_iniciais=dados)
    doc = Document(io.BytesIO(docx))
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                if "Segmento da empresa" in celula.text:
                    assert MARCADOR_PENDENTE in celula.text


def test_prefill_de_pf_nao_quebra() -> None:
    """Transição pressupõe empresa, mas a tela deixa consultar CPF. Melhor
    aproveitar o contato do que estourar."""
    dados = dados_de_contratante(PF)
    assert dados["responsavel"] == "Vinicius Almeida"
    assert dados["cpf_responsavel"] == "111.444.777-35"
    assert "razao_social" not in dados
def test_e_um_docx_valido_e_abre() -> None:
    docx = gerar_formulario_transicao(dados_iniciais=dados_de_contratante(PJ))
    assert docx[:2] == b"PK"  # zip
    assert len(Document(io.BytesIO(docx)).tables) > 5


def test_chave_desconhecida_falha_alto_em_vez_de_sumir() -> None:
    """Chave errada tem de estourar, não ser ignorada.

    A primeira versão deste teste afirmava o contrário — "melhor ignorar do que
    derrubar a geração". Está errado, e a razão é o ``.get(chave, "")`` das
    seções: uma chave com um caractere a menos não estoura, devolve vazio, e o
    campo sai como ``[ PREENCHER AQUI ]``. O documento chega ao cliente
    perguntando algo que a equipe já tinha respondido, e ninguém descobre,
    porque campo em branco por engano é idêntico a campo em branco de
    propósito.

    É a mesma falha do ``valor_mesal`` no contrato, e merece a mesma resposta.
    """
    with pytest.raises(CampoDesconhecidoError) as erro:
        gerar_formulario_transicao(dados_iniciais={"campo_que_nao_existe": "x"})
    assert "campo_que_nao_existe" in str(erro.value)


def test_erro_de_digitacao_sugere_o_campo_certo() -> None:
    """Um "s" a menos é o erro provável, e a mensagem tem de encurtar a
    caçada."""
    with pytest.raises(CampoDesconhecidoError) as erro:
        gerar_formulario_transicao(dados_pessoal={"tem_funcionario": "Não"})
    assert "tem_funcionarios" in str(erro.value)


def test_todo_campo_declarado_e_de_fato_usado_pelo_documento() -> None:
    """O inverso da validação: uma chave listada em ``CAMPOS_POR_BLOCO`` que o
    exportador não lê seria aceita pela API e sumiria sem aviso."""
    from src.exporters.docx_transicao import CAMPOS_POR_BLOCO

    argumento = {
        "iniciais": "dados_iniciais", "pessoal": "dados_pessoal",
        "fiscal": "dados_fiscal", "sucessao": "dados_sucessao",
    }
    for bloco, chaves in CAMPOS_POR_BLOCO.items():
        for chave in chaves:
            marca = f"VALOR-SENTINELA-{chave.upper().replace('_', '-')}"
            docx = gerar_formulario_transicao(**{argumento[bloco]: {chave: marca}})
            assert marca in _texto(docx), (
                f"'{chave}' está declarado no bloco '{bloco}' mas não aparece "
                f"no documento gerado"
            )
