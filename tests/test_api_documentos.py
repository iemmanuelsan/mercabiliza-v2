"""Testes da API de documentos.

O foco não é "a rota responde 200" — é provar que o documento que sai pela API
é o MESMO que sai pela aba Streamlit. Se um dia divergirem, é bug: as duas
pontas chamam as mesmas funções de propósito, e ter duas versões do contrato
circulando seria o pior defeito possível neste sistema.
"""

from __future__ import annotations

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient

import api as modulo_api

CHAVE = "chave-de-teste-longa-o-suficiente"

PJ = {
    "tipo": "PJ",
    "razao_social": "MERCADO TESTE LTDA",
    "nome_fantasia": "Mercado Teste",
    "cnpj": "11222333000181",
    "cnae_principal": "4712-1/00",
    "endereco": {
        "logradouro": "RUA ANCHIETA",
        "numero": "204",
        "bairro": "VILA BOAVENTURA",
        "municipio": "Jundiaí",
        "uf": "SP",
        "cep": "13201804",
    },
    "telefone": "(19) 3327-0038",
    "email": "contato@mercadoteste.com.br",
    "representante": {
        "nome": "Fulano de Tal",
        "cpf": "11144477735",
        "rg": "123456789",
        "orgao_emissor": "SSP/SP",
        "estado_civil": "casado",
        "profissao": "empresário",
        "qualificacao": "sócio administrador",
    },
}

PF = {
    "tipo": "PF",
    "nome": "Maria de Souza",
    "cpf": "11144477735",
    "rg": "987654321",
    "orgao_emissor": "SSP/SP",
    "estado_civil": "solteira",
    "profissao": "comerciante",
    "genero_feminino": True,
    "endereco": {
        "logradouro": "RUA DAS FLORES",
        "numero": "50",
        "bairro": "CENTRO",
        "municipio": "Campinas",
        "uf": "SP",
        "cep": "13010000",
    },
}


@pytest.fixture
def cliente(monkeypatch):
    monkeypatch.setenv("DOCUMENTOS_API_KEY", CHAVE)
    return TestClient(modulo_api.app)


@pytest.fixture
def autorizado(cliente):
    cliente.headers.update({"X-API-Key": CHAVE})
    return cliente


# --------------------------------------------------------------------------- #
# Autenticação — o que impede uma fábrica aberta de documentos timbrados       #
# --------------------------------------------------------------------------- #
def test_sem_chave_recusa(cliente):
    r = cliente.post("/v1/ficha", json={"contratante": PJ})
    assert r.status_code == 401


def test_chave_errada_recusa(cliente):
    r = cliente.post("/v1/ficha", json={"contratante": PJ}, headers={"X-API-Key": "chute"})
    assert r.status_code == 401


def test_chave_nao_configurada_falha_FECHADA(monkeypatch):
    """Sem configuração, a API não atende ninguém.

    O contrário — abrir por falta de configuração — seria a pior falha
    possível num serviço que estampa o timbrado da empresa.
    """
    monkeypatch.delenv("DOCUMENTOS_API_KEY", raising=False)
    c = TestClient(modulo_api.app)
    r = c.post("/v1/ficha", json={"contratante": PJ}, headers={"X-API-Key": "qualquer"})
    assert r.status_code == 503


def test_health_nao_exige_chave_e_denuncia_falta_de_config(monkeypatch):
    monkeypatch.delenv("DOCUMENTOS_API_KEY", raising=False)
    c = TestClient(modulo_api.app)
    r = c.get("/health")
    assert r.status_code == 200
    assert r.json()["chave_configurada"] is False


# --------------------------------------------------------------------------- #
# Ficha cadastral                                                             #
# --------------------------------------------------------------------------- #
def test_ficha_pj_devolve_pdf(autorizado):
    r = autorizado.post("/v1/ficha", json={"contratante": PJ})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert len(r.content) > 5_000


def test_ficha_pf_devolve_pdf(autorizado):
    r = autorizado.post("/v1/ficha", json={"contratante": PF})
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF")


def test_nome_do_arquivo_usa_o_cliente(autorizado):
    """O cliente recebe "ficha-cadastral-mercado-teste-ltda.pdf".

    Nome genérico ("download.pdf") vira confusão na caixa de entrada de quem
    recebe três documentos no mesmo e-mail.
    """
    r = autorizado.post("/v1/ficha", json={"contratante": PJ})
    assert "mercado-teste-ltda" in r.headers["content-disposition"]
    assert r.headers["content-disposition"].startswith("attachment")


def test_acento_no_nome_do_arquivo_e_normalizado(autorizado):
    pj = {**PJ, "razao_social": "PADARIA AÇÃO & CIA LTDA"}
    r = autorizado.post("/v1/ficha", json={"contratante": pj})
    assert r.status_code == 200
    disposicao = r.headers["content-disposition"]
    # Cabeçalho HTTP com acento quebra em parte dos clientes.
    assert disposicao.isascii(), disposicao
    assert "padaria-acao" in disposicao


# --------------------------------------------------------------------------- #
# Contrato                                                                    #
# --------------------------------------------------------------------------- #
def test_contrato_pj_devolve_pdf(autorizado):
    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": PJ,
            "parametros": {"valor_mensal": 350.0, "foro": "Campinas/SP"},
        },
    )
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF")
    assert len(r.content) > 10_000


def test_contrato_pf_devolve_pdf(autorizado):
    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": PF,
            "parametros": {"valor_mensal": 350.0},
        },
    )
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF")


def test_contrato_sem_parametros_usa_os_padroes(autorizado):
    """Não pode exigir preencher 17 campos para ver uma minuta."""
    r = autorizado.post("/v1/contrato", json={"contratante": PJ})
    assert r.status_code == 200, r.text


def test_contrato_traz_os_dados_do_cliente_no_texto(autorizado):
    """Um PDF que gera sem erro mas sai em branco passaria nos testes acima."""
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": PJ,
            "parametros": {"valor_mensal": 350.0, "foro": "Campinas/SP"},
        },
    )
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)

    assert "MERCADO TESTE LTDA" in texto
    assert "11.222.333/0001-81" in texto
    assert "Fulano de Tal" in texto
    # A CONTRATADA é texto fixo definido pela diretoria. Desde a revisão de
    # 01/09/2026 é uma empresa só — a de tecnologia saiu do contrato.
    assert "MERCABILIZA SOLUCOES FISCAIS E CONTABEIS LTDA" in texto
    assert "MERCABILIZA SOLUCOES EM TECNOLOGIA E GESTAO" not in texto
    assert "LUIS FELIPE PEDI SILVA" in texto
    assert "SP323775/O-2" in texto


def test_acentuacao_sai_correta_no_pdf(autorizado):
    """Regressão: sem a fonte DejaVu no container, "Jundiaí" saía "Jundia"."""
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    r = autorizado.post("/v1/contrato", json={"contratante": PJ})
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)
    assert "Jundiaí" in texto


# --------------------------------------------------------------------------- #
# Formulário DOCX                                                             #
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("perfil", ["PJ", "MEI", "PF"])
def test_formulario_nos_tres_perfis(autorizado, perfil):
    contratante = PF if perfil == "PF" else PJ
    corpo: dict = {"contratante": contratante, "perfil": perfil}
    if perfil == "MEI":
        corpo["desenquadramento"] = {
            "cnpj": "11222333000181",
            "razao_atual": "MERCADO TESTE LTDA",
            "abertura": "01/02/2020",
        }

    r = autorizado.post("/v1/formulario", json=corpo)
    assert r.status_code == 200, r.text
    # DOCX é um zip: começa com "PK".
    assert r.content.startswith(b"PK")
    assert "wordprocessingml" in r.headers["content-type"]
    assert f"formulario-{perfil.lower()}" in r.headers["content-disposition"]


def test_formulario_pj_vem_com_dados_preenchidos(autorizado):
    """O preenchimento híbrido: o que o sistema sabe já vem no documento."""
    docx = pytest.importorskip("docx")
    import io

    r = autorizado.post("/v1/formulario", json={"contratante": PJ, "perfil": "PJ"})
    documento = docx.Document(io.BytesIO(r.content))
    texto = "\n".join(
        celula.text
        for tabela in documento.tables
        for linha in tabela.rows
        for celula in linha.cells
    )
    assert "MERCADO TESTE LTDA" in texto


def test_formulario_pf_deixa_a_empresa_em_branco(autorizado):
    """Na abertura não existe empresa ainda: o bloco é decisão do cliente.

    Se viesse preenchido, o cliente assinaria uma razão social que o sistema
    inventou.
    """
    docx = pytest.importorskip("docx")
    import io

    r = autorizado.post("/v1/formulario", json={"contratante": PF, "perfil": "PF"})
    documento = docx.Document(io.BytesIO(r.content))
    texto = "\n".join(
        celula.text
        for tabela in documento.tables
        for linha in tabela.rows
        for celula in linha.cells
    )
    assert "PREENCHER" in texto.upper()


# --------------------------------------------------------------------------- #
# Validação de entrada                                                        #
# --------------------------------------------------------------------------- #
def test_contratante_incompleto_da_422_e_nao_500(autorizado):
    """Erro do pedido tem de ser 4xx: 500 faria o front culpar o servidor."""
    r = autorizado.post("/v1/ficha", json={"contratante": {"tipo": "PJ"}})
    assert r.status_code == 422


def test_campo_desconhecido_da_422_e_nao_e_ignorado(autorizado):
    """Nome de campo digitado errado tem de estourar, não virar o padrão.

    Se `valor_mesal` fosse ignorado em silêncio, o contrato sairia com o valor
    padrão e ninguém descobriria até o cliente questionar a cobrança. É a
    mesma classe de falha do `incluir_dp`.
    """
    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": PJ,
            "parametros": {"valor_mesal": 350.0},  # 'mensal' escrito errado
        },
    )
    assert r.status_code == 422
    assert "valor_mesal" in r.text


def test_campo_desconhecido_no_contratante_tambem_estoura(autorizado):
    r = autorizado.post(
        "/v1/ficha",
        json={
            "contratante": {**PJ, "campo_inventado": "x"},
        },
    )
    assert r.status_code == 422


def test_perfil_invalido_da_422(autorizado):
    r = autorizado.post("/v1/formulario", json={"contratante": PJ, "perfil": "XPTO"})
    assert r.status_code == 422


# --------------------------------------------------------------------------- #
# A trava que impede a divergência de padrões de voltar                        #
# --------------------------------------------------------------------------- #
def test_schema_da_api_nao_pode_ter_padrao_proprio():
    """Nenhum campo de ``ParametrosIn`` pode ter padrão diferente de ``None``.

    Foi exatamente esse erro que fez o contrato da API sair sem a cláusula
    trabalhista e sem a de monofásico: o schema dizia ``incluir_dp = False``
    enquanto o padrão real era ``True``. Com tudo em ``None``, o dataclass
    ``ParametrosContrato`` é a única fonte da verdade.
    """
    from api import ParametrosIn

    culpados = {
        nome: campo.default
        for nome, campo in ParametrosIn.model_fields.items()
        if campo.default is not None
    }
    assert not culpados, (
        "Estes campos têm padrão próprio e vão divergir do ParametrosContrato: "
        f"{culpados}. Use None e deixe o dataclass decidir."
    )


def test_todo_campo_do_dataclass_existe_no_schema():
    """Campo novo no contrato não pode ficar inacessível pela API.

    Sem esta trava, alguém acrescenta um parâmetro em ParametrosContrato, a
    aba Streamlit passa a usá-lo, e a API continua gerando com o padrão —
    divergência silenciosa nos documentos.
    """
    import dataclasses

    from api import ParametrosIn
    from src.core.contrato import ParametrosContrato

    do_dominio = {f.name for f in dataclasses.fields(ParametrosContrato)}
    do_schema = set(ParametrosIn.model_fields)
    assert do_dominio == do_schema, (
        f"só no domínio: {do_dominio - do_schema} | só no schema: {do_schema - do_dominio}"
    )


# --------------------------------------------------------------------------- #
# Paridade com a aba Streamlit — o teste que mais importa                      #
# --------------------------------------------------------------------------- #
def test_api_e_streamlit_geram_o_MESMO_contrato(autorizado):
    """Duas versões do contrato circulando seria o pior defeito deste sistema.

    Compara o texto extraído, não os bytes: o PDF carrega data de criação e
    identificadores que mudam a cada geração.
    """
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    from api import ContratantePJIn
    from src.core.contrato import ParametrosContrato
    from src.core.pessoas import contratada_padrao
    from src.exporters.pdf_documentos import gerar_contrato

    parametros = {"valor_mensal": 350.0, "foro": "Campinas/SP"}

    via_api = autorizado.post(
        "/v1/contrato", json={"contratante": PJ, "parametros": parametros}
    ).content

    direto = gerar_contrato(
        ContratantePJIn(**PJ).para_dominio(),
        contratada_padrao(),
        ParametrosContrato(valor_mensal=350.0, foro="Campinas/SP"),
    )

    def texto_de(conteudo: bytes) -> str:
        with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)

    assert texto_de(via_api) == texto_de(direto)


# --------------------------------------------------------------------------- #
# Cláusulas particulares pela API                                             #
# --------------------------------------------------------------------------- #
def test_clausulas_particulares_entram_no_contrato(autorizado):
    """O texto digitado na tela tem de aparecer no PDF, antes do foro."""
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    clausula = "Os três primeiros meses terão desconto de 20% sobre os honorários."
    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": PJ,
            "parametros": {"valor_mensal": 350.0, "clausulas_particulares": [clausula]},
        },
    )
    assert r.status_code == 200, r.text

    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)

    assert "DISPOSIÇÕES PARTICULARES" in texto
    assert "desconto de 20%" in texto
    # E o foro é renumerado: sem cláusula é a 7, com cláusula vira a 8.
    assert "CLÁUSULA 8 - DO FORO" in texto


def test_sem_clausulas_o_foro_continua_sendo_a_7(autorizado):
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    r = autorizado.post("/v1/contrato", json={"contratante": PJ})
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)

    assert "DISPOSIÇÕES PARTICULARES" not in texto
    assert "CLÁUSULA 7 - DO FORO" in texto


def test_varias_clausulas_saem_todas(autorizado):
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": PJ,
            "parametros": {
                "clausulas_particulares": [
                    "Primeira condição especial acordada entre as partes.",
                    "Segunda condição especial acordada entre as partes.",
                    "Terceira condição especial acordada entre as partes.",
                ]
            },
        },
    )
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)

    for ordinal in ("Primeira", "Segunda", "Terceira"):
        assert f"{ordinal} condição especial" in texto


def test_inscricao_estadual_entra_na_qualificacao(autorizado):
    """A IE consultada no dossiê tem de chegar ao texto do contrato."""
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    r = autorizado.post(
        "/v1/contrato",
        json={
            "contratante": {**PJ, "inscricao_estadual": "159384180119"},
        },
    )
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)

    assert "inscrição estadual nº 159384180119" in texto


def test_complemento_sujo_da_receita_sai_limpo_pela_api(autorizado):
    """Fecha o ciclo do defeito relatado: "palavra ;palavra" e "casa casa"."""
    pdfplumber = pytest.importorskip("pdfplumber")
    import io

    sujo = {
        **PJ,
        "endereco": {
            **PJ["endereco"],
            "complemento": "CASA CASA CASA CASA ;CASA TERREO ;CASA TERREO",
        },
    }
    r = autorizado.post("/v1/ficha", json={"contratante": sujo})
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        texto = "\n".join((p.extract_text() or "") for p in pdf.pages)

    assert " ;" not in texto
    assert "CASA CASA" not in texto


# --------------------------------------------------------------------------- #
# /v1/transicao — formulário de entrada por transição contábil                 #
# --------------------------------------------------------------------------- #
def _texto_docx(conteudo: bytes) -> str:
    import io

    from docx import Document

    doc = Document(io.BytesIO(conteudo))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                partes.append(celula.text)
    return "\n".join(partes)


def test_transicao_exige_chave(cliente):
    assert cliente.post("/v1/transicao", json={}).status_code == 401


def test_transicao_sem_nenhum_dado_gera_o_formulario_em_branco(autorizado):
    """É o impresso que a equipe usava na planilha. Tem de sair sem exigir
    consulta de CNPJ nenhuma."""
    r = autorizado.post("/v1/transicao", json={})
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    assert "FORMULÁRIO DE ENTRADA DE NOVOS CLIENTES" in _texto_docx(r.content)


def test_transicao_com_contratante_preenche_o_cadastral(autorizado):
    r = autorizado.post("/v1/transicao", json={"contratante": PJ})
    assert r.status_code == 200
    texto = _texto_docx(r.content)
    assert "MERCADO TESTE LTDA" in texto
    assert "11.222.333/0001-81" in texto


def test_transicao_leva_as_respostas_da_tela(autorizado):
    r = autorizado.post(
        "/v1/transicao",
        json={
            "contratante": PJ,
            "iniciais": {"competencia": "09/2026", "segmento": "MiniMercado Autônomo"},
            "pessoal": {"tem_funcionarios": "Não"},
            "fiscal": {"sistema_notas": "Bling"},
            "sucessao": {"email_anterior": "contato@agilize.com.br"},
        },
    )
    assert r.status_code == 200
    texto = _texto_docx(r.content)
    for esperado in ("09/2026", "MiniMercado Autônomo", "Bling", "contato@agilize.com.br"):
        assert esperado in texto


def test_transicao_a_tela_vence_o_prefill(autorizado):
    """Se a equipe corrigiu um dado na tela, a correção manda. O contrário
    faria a consulta de CNPJ sobrescrever a conferência humana — que é
    justamente o que o documento existe para registrar."""
    r = autorizado.post(
        "/v1/transicao",
        json={
            "contratante": PJ,
            "iniciais": {"razao_social": "NOME CORRIGIDO NA JUNTA LTDA"},
        },
    )
    texto = _texto_docx(r.content)
    assert "NOME CORRIGIDO NA JUNTA LTDA" in texto
    assert "MERCADO TESTE LTDA" not in texto


def test_transicao_campo_digitado_errado_e_422_e_nao_documento_em_branco(autorizado):
    """A falha do ``valor_mesal``, de novo. Sem ``extra="forbid"``, um "s" a
    menos em ``tem_funcionarios`` geraria um documento com o campo em branco e
    resposta 200 — indistinguível de um formulário legitimamente não
    preenchido."""
    r = autorizado.post(
        "/v1/transicao",
        json={
            "pessoal": {"tem_funcionario": "Não"},
        },
    )
    assert r.status_code == 422
    assert "tem_funcionario" in r.text


def test_schema_da_api_e_exportador_nao_podem_divergir():
    """Os nomes dos campos da API têm de ser exatamente as chaves que o
    exportador conhece.

    Sem este teste, acrescentar um campo em um lado só produz um campo que a
    tela envia e o documento ignora — silenciosamente, porque o exportador lê
    com ``.get`` e a API aceitaria o extra.
    """
    from src.exporters.docx_transicao import CAMPOS_POR_BLOCO

    modelos = {
        "iniciais": modulo_api.IniciaisIn,
        "pessoal": modulo_api.PessoalIn,
        "fiscal": modulo_api.FiscalIn,
        "sucessao": modulo_api.SucessaoIn,
    }
    assert set(modelos) == set(CAMPOS_POR_BLOCO)
    for bloco, modelo in modelos.items():
        assert set(modelo.model_fields) == set(CAMPOS_POR_BLOCO[bloco]), bloco


def test_schema_da_transicao_nao_pode_ter_padrao_proprio():
    """Nenhum campo dos blocos pode ter valor padrão.

    O padrão mora no exportador. Repetido aqui, os dois divergem — foi
    exatamente assim que o ``incluir_dp`` saiu ``False`` na API e ``True`` no
    Streamlit, e os contratos gerados pela web saíram sem a cláusula
    trabalhista.
    """
    for modelo in (
        modulo_api.IniciaisIn,
        modulo_api.PessoalIn,
        modulo_api.FiscalIn,
        modulo_api.SucessaoIn,
    ):
        for nome, campo in modelo.model_fields.items():
            assert campo.default is None, (
                f"{modelo.__name__}.{nome} tem padrão próprio "
                f"({campo.default!r}); o padrão pertence ao exportador."
            )


def test_transicao_api_bate_com_o_exportador_direto(autorizado):
    """Paridade API × Streamlit: os dois caminhos têm de produzir o mesmo
    documento. Comparo o TEXTO, porque os bytes carregam a data de geração e
    ids aleatórios do pacote OOXML."""
    from src.core.pessoas import ContratantePJ
    from src.exporters.docx_transicao import (
        dados_de_contratante,
        gerar_formulario_transicao,
    )

    corpo = {"contratante": PJ, "pessoal": {"tem_funcionarios": "Não"}}
    pela_api = autorizado.post("/v1/transicao", json=corpo)

    contratante = modulo_api.ContratantePJIn(**PJ).para_dominio()
    assert isinstance(contratante, ContratantePJ)
    direto = gerar_formulario_transicao(
        dados_iniciais=dados_de_contratante(contratante),
        dados_pessoal={"tem_funcionarios": "Não"},
    )

    assert _texto_docx(pela_api.content) == _texto_docx(direto)


def test_bloco_removido_nao_e_aceito_em_silencio(autorizado):
    """Os blocos que o Iago cortou (onboarding, contábil, financeiro,
    observações) não existem mais no schema.

    Se alguém mandar um deles — uma tela desatualizada, um script antigo — o
    ``extra="forbid"`` responde 422 nomeando o bloco. Sem isso, o corpo seria
    aceito, o bloco descartado, e o documento sairia sem aquelas respostas sem
    nenhum aviso.
    """
    for bloco in ("onboarding", "contabil", "financeiro", "observacoes"):
        r = autorizado.post("/v1/transicao", json={bloco: {"x": "y"}})
        assert r.status_code == 422, bloco
        assert bloco in r.text
<<<<<<< Updated upstream
=======


# --------------------------------------------------------------------------- #
# /v1/cartao-cnpj                                                              #
# --------------------------------------------------------------------------- #
CARTAO = {
    "cnpj": "62350925000110",
    "razao_social": "MERCABILIZA SOLUCOES FISCAIS E CONTABEIS LTDA",
    "nome_fantasia": "MERCABILIZA",
    "data_abertura": "2025-08-22",
    "natureza_juridica": "206-2 - Sociedade Empresária Limitada",
    "porte": "DEMAIS",
    "email": "contato@mercabiliza.com.br",
    "telefone": "1933270038",
    "situacao": "ATIVA",
    "data_situacao": "2025-08-22",
    "endereco": {
        "logradouro": "RUA ANCHIETA",
        "numero": "204",
        "complemento": "SALA 102",
        "bairro": "VILA BOAVENTURA",
        "municipio": "JUNDIAI",
        "uf": "SP",
        "cep": "13201804",
    },
    "atividade_principal": {"codigo": "6920601", "descricao": "Atividades de contabilidade"},
    "atividades_secundarias": [
        {"codigo": "8211300", "descricao": "Serviços combinados de escritório"}
    ],
    "fontes": ["BrasilAPI", "CNPJ.ws"],
}


def test_cartao_exige_chave(cliente):
    assert cliente.post("/v1/cartao-cnpj", json=CARTAO).status_code == 401


def test_cartao_sai_em_pdf_com_nome_util(autorizado):
    r = autorizado.post("/v1/cartao-cnpj", json=CARTAO)

    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    # O nome do arquivo é o que o contador vai achar na pasta de Downloads
    # daqui a três semanas.
    assert "cartao-cnpj-mercabiliza" in r.headers["content-disposition"]


def test_cartao_so_precisa_do_cnpj(autorizado):
    """CNPJ recém-aberto vem com metade dos campos em branco nas APIs
    públicas. O comprovante tem que sair mesmo assim."""
    r = autorizado.post("/v1/cartao-cnpj", json={"cnpj": "62350925000110"})
    assert r.status_code == 200
    assert len(r.content) > 1000


def test_cartao_recusa_campo_digitado_errado(autorizado):
    """`extra="forbid"`: campo com nome errado vira 422 na hora, em vez de um
    documento com o campo em branco que ninguém percebe."""
    r = autorizado.post("/v1/cartao-cnpj", json={**CARTAO, "razao_sozial": "X"})
    assert r.status_code == 422


def test_cartao_nao_aceita_diagnostico_vindo_do_navegador(autorizado):
    """A classificação tributária do CNAE é calculada no servidor.

    Aceitá-la pronta deixaria o navegador decidir anexo e alíquota — e esse
    número vira proposta comercial.
    """
    r = autorizado.post(
        "/v1/cartao-cnpj",
        json={
            **CARTAO,
            "atividade_principal": {
                "codigo": "6920601",
                "descricao": "Atividades de contabilidade",
                "diagnostico": {"anexo": "I", "aliquota": 0.04},
            },
        },
    )
    assert r.status_code == 422
>>>>>>> Stashed changes
