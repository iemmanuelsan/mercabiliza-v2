"""Motor compartilhado dos formulários DOCX.

Nasceu de uma extração: o formulário de abertura/desenquadramento já tinha
todo este maquinário, e o de transição contábil precisava do mesmo. Duplicar
significaria que uma correção de layout entraria em um e não no outro — e a
divergência só apareceria quando um cliente recebesse dois documentos da mesma
contabilidade com cara diferente.

## A convenção visual, que é o ponto de tudo isto

* **Negrito** = preenchido pelo sistema (API de CNPJ, CEP ou cadastro).
  O cliente só confere.
* ``[ PREENCHER AQUI ]`` em vermelho, sobre fundo rosado = depende de decisão
  ou informação que só o cliente tem.
* Linha pontilhada = campo opcional, pode ficar em branco.

Sem essa distinção o cliente relê o documento inteiro. Com ela, ele vai direto
ao que falta — que costuma ser um punhado de campos.

## Por que DOCX e não XLSX

O formulário é um documento com campos rotulados, não uma planilha de dados.
Em DOCX o cliente preenche no Word, no Google Docs ou imprime e escreve à mão,
sem quebrar layout. XLSX faria sentido se o retorno fosse importado de volta em
massa; para um formulário que o cliente assina, DOCX é o certo.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..config import (
    CONTRATADA_EMAIL,
    CONTRATADA_NOME_FANTASIA,
    CONTRATADA_TELEFONE,
)

MARCADOR_PENDENTE = "[ PREENCHER AQUI ]"

COR_MARCA = RGBColor(0xDC, 0x32, 0x50)
COR_PENDENTE = RGBColor(0xC0, 0x1B, 0x36)
COR_CINZA = RGBColor(0x55, 0x5B, 0x66)
COR_BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CINZA_FUNDO = "F5F6F8"
AZUL_SECAO = "2F5597"
ROSA_PENDENTE = "FDF0F2"


# --------------------------------------------------------------------------- #
# Modelo de campo                                                              #
# --------------------------------------------------------------------------- #
@dataclass(frozen=True, slots=True)
class Campo:
    """Um campo do formulário.

    ``valor`` vazio + ``pendente=True`` → sai como ``[ PREENCHER AQUI ]``.
    ``valor`` preenchido → sai em negrito (veio do sistema).
    ``valor`` vazio + ``pendente=False`` → sai como linha em branco.
    """

    rotulo: str
    valor: str = ""
    pendente: bool = False
    dica: str = ""

    @property
    def preenchido_pelo_sistema(self) -> bool:
        return bool(self.valor.strip())


@dataclass(slots=True)
class Secao:
    titulo: str
    campos: list[Campo] = field(default_factory=list)
    colunas: int = 2


# --------------------------------------------------------------------------- #
# Helpers de formatação                                                        #
# --------------------------------------------------------------------------- #
def sombrear(celula, cor_hex: str) -> None:
    """Fundo de célula — o python-docx não expõe isso na API pública."""
    elemento = OxmlElement("w:shd")
    elemento.set(qn("w:val"), "clear")
    elemento.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(elemento)


def run(paragrafo, texto: str, *, negrito: bool = False, tamanho: float = 9,
        cor: RGBColor | None = None, italico: bool = False):
    r = paragrafo.add_run(texto)
    r.bold = negrito
    r.italic = italico
    r.font.size = Pt(tamanho)
    if cor is not None:
        r.font.color.rgb = cor
    return r


def escrever_campo(celula, campo: Campo) -> None:
    """Rótulo + valor numa célula, aplicando a convenção visual."""
    celula.text = ""
    p = celula.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, f"{campo.rotulo}: ", negrito=False, tamanho=8.5, cor=COR_CINZA)

    if campo.preenchido_pelo_sistema:
        run(p, campo.valor, negrito=True, tamanho=9)
    elif campo.pendente:
        run(p, MARCADOR_PENDENTE, negrito=True, tamanho=9, cor=COR_PENDENTE)
    else:
        run(p, "_" * 28, tamanho=9, cor=COR_CINZA)

    if campo.dica:
        p2 = celula.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        run(p2, campo.dica, tamanho=7, cor=COR_CINZA, italico=True)


def titulo_secao(doc, texto: str) -> None:
    tabela = doc.add_table(rows=1, cols=1)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    celula = tabela.rows[0].cells[0]
    sombrear(celula, AZUL_SECAO)
    celula.text = ""
    p = celula.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, texto.upper(), negrito=True, tamanho=9.5, cor=COR_BRANCO)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def tabela_campos(doc, campos: list[Campo], colunas: int = 2) -> None:
    if not campos:
        return
    linhas = -(-len(campos) // colunas)
    tabela = doc.add_table(rows=linhas, cols=colunas)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, campo in enumerate(campos):
        celula = tabela.cell(i // colunas, i % colunas)
        escrever_campo(celula, campo)
        if campo.pendente and not campo.preenchido_pelo_sistema:
            sombrear(celula, ROSA_PENDENTE)

    # Células sobrando na última linha ficam em branco, não com "None".
    for j in range(len(campos), linhas * colunas):
        tabela.cell(j // colunas, j % colunas).text = ""

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _borda_inferior(paragrafo) -> None:
    """Régua de escrita como BORDA do parágrafo, não como underscores.

    A primeira versão usava ``"_" * 95``. O problema é que o comprimento de uma
    régua de underscores depende da fonte: no render ela ficava com ~55% da
    largura útil e o bloco parecia desalinhado do resto do documento, que usa
    tabelas de ponta a ponta. Borda de parágrafo acompanha a margem sozinha,
    em qualquer fonte e em qualquer papel.
    """
    # ⚠️ O ``w:between`` é obrigatório e não é detalhe estético.
    #
    # O Word (e o LibreOffice) tratam parágrafos CONSECUTIVOS com bordas
    # IDÊNTICAS como um bloco único e desenham só a moldura externa. Com
    # apenas ``w:bottom``, quatro linhas de resposta viravam UMA — e o defeito
    # não aparece no XML nem no texto extraído, só no render. O ``w:between``
    # é o que manda desenhar a divisória entre os parágrafos do grupo.
    pPr = paragrafo._p.get_or_add_pPr()
    bordas = OxmlElement("w:pBdr")
    for lado in ("bottom", "between"):
        borda = OxmlElement(f"w:{lado}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "6")
        borda.set(qn("w:space"), "1")
        borda.set(qn("w:color"), "AAB0BA")
        bordas.append(borda)
    pPr.append(bordas)


def linhas_para_escrever(doc, quantidade: int = 4) -> None:
    """Linhas em branco para resposta discursiva.

    O Excel resolvia isso com bordas inferiores em células vazias; aqui é a
    mesma ideia, e sobrevive à impressão, ao Google Docs e ao preenchimento à
    mão.
    """
    for _ in range(quantidade):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        _borda_inferior(p)


# --------------------------------------------------------------------------- #
# Estrutura do documento                                                       #
# --------------------------------------------------------------------------- #
def configurar_pagina(doc) -> None:
    for secao in doc.sections:
        secao.top_margin = Cm(1.8)
        secao.bottom_margin = Cm(1.8)
        secao.left_margin = Cm(1.8)
        secao.right_margin = Cm(1.8)


def cabecalho(doc, titulo: str, subtitulo: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, CONTRATADA_NOME_FANTASIA.upper(), negrito=True, tamanho=16,
        cor=COR_MARCA)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, titulo, negrito=True, tamanho=13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, subtitulo, tamanho=8.5, cor=COR_CINZA)


def legenda(doc) -> None:
    tabela = doc.add_table(rows=1, cols=1)
    tabela.style = "Table Grid"
    celula = tabela.rows[0].cells[0]
    sombrear(celula, CINZA_FUNDO)
    celula.text = ""
    p = celula.paragraphs[0]
    run(p, "Como preencher: ", negrito=True, tamanho=8.5)
    run(p, "os campos em ", tamanho=8.5)
    run(p, "negrito", negrito=True, tamanho=8.5)
    run(p, " já foram preenchidos pela contabilidade — apenas confira. Os "
           "campos marcados como ", tamanho=8.5)
    run(p, MARCADOR_PENDENTE, negrito=True, tamanho=8.5, cor=COR_PENDENTE)
    run(p, " dependem da sua decisão. Os demais são dados que você deve "
           "informar.", tamanho=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def rodape_declaracao(doc, declaracao: str) -> None:
    """Só o texto da declaração — sem assinatura, sem local/data, sem contatos.

    ⚠️ **Isto não é o mesmo que ``rodape_assinatura`` com menos coisas.**

    Sem linha de assinatura, o documento deixa de ser um termo e passa a ser um
    formulário de coleta: o texto informa o cliente do que será feito com os
    dados, mas ninguém atesta nada nem autoriza formalmente. Qualquer trecho que
    afirme "a assinatura abaixo concede a autorização" vira mentira aqui — a
    redação tem de acompanhar.

    Os contatos da CONTRATADA também ficam de fora de propósito: os valores em
    ``config.py`` são de outra empresa, e imprimir contato errado é pior do que
    não imprimir contato nenhum.
    """
    doc.add_paragraph()
    p = doc.add_paragraph()
    run(p, "Declaração: ", negrito=True, tamanho=8)
    run(p, declaracao, tamanho=8)


def rodape_assinatura(doc, declaracao: str, rotulo_assinatura: str) -> None:
    doc.add_paragraph()
    p = doc.add_paragraph()
    run(p, "Declaração: ", negrito=True, tamanho=8)
    run(p, declaracao, tamanho=8)

    doc.add_paragraph()
    tabela = doc.add_table(rows=2, cols=2)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, rotulo in enumerate([rotulo_assinatura, "Local e data"]):
        c = tabela.cell(0, col)
        c.text = ""
        run(c.paragraphs[0], "_" * 40, tamanho=10)
        c2 = tabela.cell(1, col)
        c2.text = ""
        run(c2.paragraphs[0], rotulo, tamanho=8, cor=COR_CINZA)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, f"{CONTRATADA_NOME_FANTASIA} · {CONTRATADA_TELEFONE} · "
           f"{CONTRATADA_EMAIL}", tamanho=7.5, cor=COR_CINZA)
